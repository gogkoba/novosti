"""Собирает Word-файл из дайджеста в markdown.

Запуск:
    python tools/md2docx.py output/2026-09-09/Новости_2026-09-09.md
    python tools/md2docx.py output/2026-09-09/Новости_2026-09-09.md -o "C:/Users/.../Downloads/Новости.docx"
    python tools/md2docx.py <файл.md> --no-images      # без скачивания фото

Понимает и чистый файл (Новости_<дата>.md), и рабочий (01_digest.md): теги в квадратных
скобках и строки «одобрено» пропускает. Фото скачивает по ссылке и вставляет в документ;
если скачать не удалось — оставляет ссылку. Ссылки на источники и видео делает кликабельными.
"""

import argparse
import io
import re
import sys
import time
import urllib.request
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
BODY_PT = 12
HEAD_PT = 13
TITLE_PT = 16
SMALL_PT = 9
GREY = RGBColor(0x59, 0x59, 0x59)
LINK = RGBColor(0x1F, 0x4E, 0x79)
IMAGE_WIDTH_CM = 14
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/120 Safari/537.36"

URL_RE = re.compile(r"https?://\S+")


# ---------- разбор markdown ----------

def parse(md_text):
    """Возвращает (заголовок документа, список новостей)."""
    lines = md_text.splitlines()
    title = "Новости"
    for line in lines:
        if line.startswith("# "):
            m = re.search(r"(\d{4})-(\d{2})-(\d{2})", line)
            title = f"Новости на {m.group(3)}.{m.group(2)}.{m.group(1)}" if m else line[2:].strip()
            break

    items, cur = [], None
    for line in lines:
        if line.startswith("## "):
            if cur:
                items.append(cur)
            head = line[3:].strip()
            head = re.sub(r"^\d+\.\s*", "", head)               # свой номер: нумеруем заново
            head = re.sub(r"\s*·\s*\[[^\]]*\]\s*$", "", head)  # тег темы
            cur = {"head": head, "paras": [], "note": None, "sources": [], "photo": None, "video": None}
            continue
        if cur is None or not line.strip():
            continue
        s = line.strip()
        if s.startswith("- [") and "одобрено" in s:
            continue
        if s.startswith("Не позиция канала"):
            cur["note"] = s
        elif s.startswith("Источник:"):
            cur["sources"] = parse_links(s[len("Источник:"):])
        elif s.startswith("Фото:"):
            urls = URL_RE.findall(s)
            cur["photo"] = urls[0] if urls else None
        elif s.startswith("Видео:"):
            body = s[len("Видео:"):].strip()
            cur["video"] = None if body.lower().startswith("нет") else parse_links(body)
        else:
            cur["paras"].append(re.sub(r"\*\*(.+?)\*\*", r"\1", s))
    if cur:
        items.append(cur)
    return title, items


def parse_links(text):
    """'URL · доп.: URL · стенограмма: URL' -> [(подпись, url), ...]."""
    out, seen = [], {}
    for token in re.split(r"\s+·\s+", text.strip()):
        m = URL_RE.search(token)
        if not m:
            continue
        url = m.group(0).rstrip(".,;)")
        label = token[: m.start()].strip(" :—-–")
        if label.lower() in ("", "доп.", "доп"):
            label = host(url)
            seen[label] = seen.get(label, 0) + 1
            if seen[label] > 1:                      # второй и далее источник с того же сайта
                label = f"{label} ({seen[label]})"
        out.append((label, url))
    return out


def host(url):
    h = urlparse(url).netloc
    return h[4:] if h.startswith("www.") else h


# ---------- сборка docx ----------

def add_hyperlink(paragraph, url, text, size=SMALL_PT):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), FONT)
    rpr.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(size * 2))
    rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def small_run(paragraph, text):
    run = paragraph.add_run(text)
    run.font.size = Pt(SMALL_PT)
    run.font.color.rgb = GREY
    return run


def fetch_image(url, timeout=25, attempts=3):
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    last_exc = None
    for attempt in range(attempts):                      # сети РИА и Коммерсанта отвечают нестабильно
        try:
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                ctype = resp.headers.get("Content-Type", "")
                data = resp.read()
            if data:
                break
            last_exc = ValueError("пустой ответ")
        except Exception as exc:  # noqa: BLE001
            last_exc = exc
        time.sleep(2 * (attempt + 1))
    else:
        raise last_exc
    is_image = (
        data[:3] == b"\xff\xd8\xff"          # JPEG
        or data[:8] == b"\x89PNG\r\n\x1a\n"  # PNG
        or data[:6] in (b"GIF87a", b"GIF89a")
        or data[:2] == b"BM"                 # BMP
    )
    if not is_image:
        raise ValueError(f"not an image (Content-Type: {ctype or 'нет'}, начало: {data[:12]!r})")
    if is_headline_card(url, data):
        raise ValueError("карточка с заголовком, а не фото")
    return normalize_image(data)


def normalize_image(data, max_width=1600):
    """Перекодирует картинку в обычный JPEG с JFIF-заголовком и ужимает по ширине.
    python-docx не распознаёт JPEG без JFIF/Exif (так отдаёт, например, РИА)."""
    try:
        from PIL import Image
    except ImportError:
        return data
    im = Image.open(io.BytesIO(data))
    if im.mode not in ("RGB", "L"):
        im = im.convert("RGB")
    if im.width > max_width:
        im = im.resize((max_width, round(im.height * max_width / im.width)))
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=88)
    return buf.getvalue()


def is_headline_card(url, data):
    """Карточки с логотипом и заголовком вместо фото — в документ не вставляем.
    РИА: любой адрес вида img.ria.ru/images/sharing/... Интерфакс: /aspimg/<id>.jpg
    размером 700×350 (для статей с фото по тому же адресу лежит нормальная картинка)."""
    if "img.ria.ru/images/sharing/" in url:
        return True
    if "interfax.ru/aspimg/" not in url:
        return False
    try:
        from PIL import Image
        return Image.open(io.BytesIO(data)).size == (700, 350)
    except Exception:  # noqa: BLE001 — нет PIL или битый файл: считаем, что это фото
        return False


def build(title, items, out_path, with_images=True, log=print):
    doc = Document()

    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(TITLE_PT)
    p.paragraph_format.space_after = Pt(14)

    for i, it in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(f"{i}. {it['head']}")
        run.bold = True
        run.font.size = Pt(HEAD_PT)

        for text in it["paras"]:
            p = doc.add_paragraph(text)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        if it["note"]:
            p = doc.add_paragraph()
            run = p.add_run(it["note"])
            run.italic = True
            run.font.size = Pt(BODY_PT - 1)
            run.font.color.rgb = GREY

        photo_done = False
        if with_images and it["photo"]:
            try:
                data = fetch_image(it["photo"])
                doc.add_picture(io.BytesIO(data), width=Cm(IMAGE_WIDTH_CM))
                doc.paragraphs[-1].paragraph_format.space_after = Pt(4)
                photo_done = True
            except Exception as exc:  # noqa: BLE001 — любая ошибка сети/формата: оставляем ссылку
                log(f"  фото {i}: не скачалось ({exc}); оставляю ссылку")

        if it["sources"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            small_run(p, "Источник: ")
            for k, (label, url) in enumerate(it["sources"]):
                if k:
                    small_run(p, " · ")
                add_hyperlink(p, url, label)

        if it["photo"] and not photo_done:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            small_run(p, "Фото: ")
            add_hyperlink(p, it["photo"], host(it["photo"]))

        if it["video"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            small_run(p, "Видео: ")
            for k, (label, url) in enumerate(it["video"]):
                if k:
                    small_run(p, " · ")
                add_hyperlink(p, url, label)

    doc.save(out_path)


def main():
    ap = argparse.ArgumentParser(description="Markdown-дайджест -> Word")
    ap.add_argument("src", help="файл .md")
    ap.add_argument("-o", "--out", help="куда сохранить .docx (по умолчанию рядом с исходником)")
    ap.add_argument("--no-images", action="store_true", help="не скачивать и не вставлять фото")
    args = ap.parse_args()

    src = Path(args.src)
    out = Path(args.out) if args.out else src.with_suffix(".docx")
    title, items = parse(src.read_text(encoding="utf-8"))
    print(f"{title}: {len(items)} новостей -> {out}")
    build(title, items, out, with_images=not args.no_images)
    print(f"готово: {out} ({out.stat().st_size // 1024} КБ)")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
