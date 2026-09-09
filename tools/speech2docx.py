"""Собирает Word-файл из речи ведущего (или статьи) — обычного текста с абзацами.

Запуск:
    python tools/speech2docx.py output/2026-09-09/02_speech.md
    python tools/speech2docx.py output/2026-09-09/03_column.md -t "Авторская статья" -o "C:/.../Downloads/Статья.docx"

Оформление под чтение с листа или суфлёра: Times New Roman 14, интервал 1,5, поля 2 см.
Заголовок берётся из ключа -t (по умолчанию «Речь ведущего») плюс дата из имени папки.
"""

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "Times New Roman"
BODY_PT = 14
TITLE_PT = 16


def build(text, title, out_path):
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.line_spacing = 1.5

    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(TITLE_PT)
    p.paragraph_format.space_after = Pt(16)

    for para in [x.strip() for x in re.split(r"\n\s*\n", text) if x.strip()]:
        para = re.sub(r"^#+\s*", "", para)              # на случай markdown-заголовка
        para = re.sub(r"\*\*(.+?)\*\*", r"\1", para)
        p = doc.add_paragraph(para)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.save(out_path)


def main():
    ap = argparse.ArgumentParser(description="Речь/статья -> Word")
    ap.add_argument("src", help="файл с текстом")
    ap.add_argument("-o", "--out", help="куда сохранить .docx (по умолчанию рядом с исходником)")
    ap.add_argument("-t", "--title", default="Речь ведущего", help="заголовок документа")
    args = ap.parse_args()

    src = Path(args.src)
    out = Path(args.out) if args.out else src.with_suffix(".docx")
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", str(src.resolve()))
    title = f"{args.title}, {m.group(3)}.{m.group(2)}.{m.group(1)}" if m else args.title
    text = src.read_text(encoding="utf-8")
    build(text, title, out)
    print(f"готово: {out} ({len(text.split())} слов)")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
